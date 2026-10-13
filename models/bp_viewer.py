"""Render the Tabax business plan workbook as one commentable artifact page (values + formulas)."""
import base64, datetime, html, sys
from openpyxl import load_workbook

XLSX, SC1, SC3, OUT = sys.argv[1:5]
wv = load_workbook(XLSX, data_only=True)
wf = load_workbook(XLSX)
E = lambda s: html.escape(str(s), quote=True)


def nv(name):
    """Value of a defined name (single cell) in the recalculated workbook."""
    sheet, cell = wf.defined_names[name].attr_text.split("!")
    return wv[sheet.strip("'")][cell.replace("$", "")].value


FX = nv("FX_EUR")
FXU = nv("FX_USD")
NNBSP = " "
MINUS = "−"


def fmt_num(v, kind, scale=1):
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        return E(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%d %b %Y")
    x = v / scale
    if kind == "pct":
        s = f"{abs(x) * 100:.1f}%"
    elif kind == "pct2":
        s = f"{abs(x) * 100:.2f}%"
    elif kind == "dec1":
        s = f"{abs(x):,.1f}".replace(",", NNBSP)
    elif kind == "dec2":
        s = f"{abs(x):,.2f}".replace(",", NNBSP)
    elif kind == "dec3":
        s = f"{abs(x):,.3f}".replace(",", NNBSP)
    elif kind == "mult":
        s = f"{abs(x):.1f}×"
    else:
        s = f"{abs(round(x)):,}".replace(",", NNBSP)
    if abs(x) < 1e-9:
        return '<span class="z">–</span>'
    return (f'<span class="neg">{MINUS}{s}</span>' if x < 0 else s)


def kind_of(numfmt):
    f = numfmt or ""
    if "0.00%" in f:
        return "pct2"
    if "%" in f:
        return "pct"
    if "0.0x" in f:
        return "mult"
    if f.startswith("#,##0.0") or f == "0.0":
        return "dec1"
    if f == "0.00":
        return "dec2"
    if f in ("0.000", "#,##0.000"):
        return "dec3"
    if "yy" in f:
        return "date"
    return "int"


def eur_fmt(x):
    """Euro amount, compact: €15.09, €850, €4.2k, €61k, €1.23M, €12.4M."""
    if not isinstance(x, (int, float)) or abs(x) < 0.005:
        return ""
    a = abs(x)
    if a >= 1e7:
        t = f"{a / 1e6:.1f}M"
    elif a >= 1e6:
        t = f"{a / 1e6:.2f}M"
    elif a >= 1e4:
        t = f"{a / 1e3:.0f}k"
    elif a >= 1e3:
        t = f"{a / 1e3:.1f}k"
    elif a >= 100:
        t = f"{a:.0f}"
    else:
        t = f"{a:.2f}"
    return f'<span class="neg">{MINUS}€{t}</span>' if x < 0 else f"€{t}"


def fcfa_fmt(x):
    """FCFA amount, compact: 9 900 FCFA, 45.9M FCFA, 2.83bn FCFA."""
    a = abs(x)
    if a >= 1e9:
        t = f"{a / 1e9:.2f}bn"
    elif a >= 1e6:
        t = f"{a / 1e6:.1f}M"
    else:
        t = f"{round(a):,}".replace(",", NNBSP)
    return f"{MINUS if x < 0 else ''}{t} FCFA"


def to_eur(x, unit):
    """Euro value of an amount expressed in `unit` (FCFA..., USD...); None if not money."""
    u = unit or ""
    if not isinstance(x, (int, float)):
        return None
    if u.startswith("FCFA") or u in ("k FCFA", "M FCFA"):
        return x / FX
    if u.startswith("USD"):
        return x * FXU / FX
    return None


def dual(x, kind, scale, unit):
    """FCFA (or USD) figure with its euro equivalent underneath."""
    main = fmt_num(x, kind, scale)
    e = to_eur(x, unit)
    if e is None or not isinstance(x, (int, float)) or abs(x) < 1e-9:
        return main
    return f'{main}<span class="eur">{eur_fmt(e)}</span>'


def formula_of(ws, coord):
    v = ws[coord].value
    return v if isinstance(v, str) and v.startswith("=") else ""


# ------------------------------------------------------------------ monthly sheets
MONTH_COLS = [5 + i for i in range(60)]


def ts_table(sheet, from_label=None, k_scale=True):
    """Monthly table; with from_label, start at that section header (inclusive)."""
    v, f = wv[sheet], wf[sheet]
    head1 = "".join(f'<th scope="col">{v.cell(4, c).value.strftime("%b %y")}</th>' for c in MONTH_COLS)
    head0 = "".join(f'<th scope="col" class="yr">Y{v.cell(5, c).value}</th>' if (c - 5) % 12 == 0 else '<th scope="col"></th>' for c in MONTH_COLS)
    rows = []
    started = from_label is None
    for r in range(7, v.max_row + 1):
        a = v.cell(r, 1).value
        if not started:
            if a != from_label:
                continue
            started = True
        e_f = f.cell(r, 5).value
        if a is None and e_f is None:
            continue
        if e_f is None:  # section header
            rows.append(f'<tr class="sec"><th scope="rowgroup" colspan="{3 + len(MONTH_COLS)}">{E(a)}</th></tr>')
            continue
        unit = v.cell(r, 2).value or ""
        kind = kind_of(f.cell(r, 5).number_format)
        scale = 1000 if (k_scale and unit == "FCFA") else 1
        unit_show = "k FCFA" if scale == 1000 else unit
        bold = bool(f.cell(r, 1).font and f.cell(r, 1).font.b)
        fm = formula_of(f, f"F{r}") or formula_of(f, f"E{r}")
        tip = f' title="Formula (month 2): {E(fm)}"' if fm else ""
        tot = dual(v.cell(r, 4).value, kind, scale, unit) if f.cell(r, 4).value is not None else ""
        cells = "".join(f"<td>{dual(v.cell(r, c).value, kind, scale, unit)}</td>" for c in MONTH_COLS)
        rows.append(f'<tr data-comment-target{" class=b" if bold else ""}><th scope="row"{tip}>{E(a)}</th><td class="u">{E(unit_show)}</td><td class="t">{tot}</td>{cells}</tr>')
    return (f'<div class="tw"><table class="ts"><thead><tr><th class="c0"></th><th class="u"></th><th class="t"></th>{head0}</tr>'
            f'<tr><th class="c0" scope="col">Line</th><th class="u" scope="col">Unit</th><th class="t" scope="col">Total / end</th>{head1}</tr></thead>'
            f'<tbody>{"".join(rows)}</tbody></table></div>')


# ------------------------------------------------------------------ annual-style sheets (label in A, values in D..H/I)
def annual_table(sheet, cols="DEFGHI", head_row=4, first=6, scale_fcfa=None, last=None, stop_label=None):
    v, f = wv[sheet], wf[sheet]
    heads = []
    for i, c in enumerate(cols):
        hv = v[f"{c}{head_row}"].value
        if isinstance(hv, (int, float)) and hv >= 2000:
            heads.append(f"Year {i + 1}<br><span class='fy'>{int(hv)}</span>")
        elif isinstance(hv, (int, float)) and c != "I":
            fy = v[f"{c}5"].value if sheet == "P&L annual" else None
            heads.append(f"Year {int(hv)}" + (f"<br><span class='fy'>{int(fy)}</span>" if fy else ""))
        else:
            heads.append(E(hv or ""))
    rows = []
    for r in range(first, (last or v.max_row) + 1):
        a = v.cell(r, 1).value
        d_f = f[f"D{r}"].value
        if stop_label and a == stop_label:
            break
        if a is None and d_f is None:
            continue
        if d_f is None and all(f[f"{c}{r}"].value is None for c in cols):
            rows.append(f'<tr class="sec"><th scope="rowgroup" colspan="{len(cols) + 2}">{E(a)}</th></tr>')
            continue
        unit = v.cell(r, 2).value or ""
        kind = kind_of(f[f"D{r}"].number_format)
        lab = str(a)
        scale = 1
        src_unit = unit
        if scale_fcfa and unit == "FCFA":
            scale = scale_fcfa["div"]
            unit = scale_fcfa["unit"]
        bold = bool(f.cell(r, 1).font and f.cell(r, 1).font.b)
        fm = formula_of(f, f"E{r}") or formula_of(f, f"D{r}")
        tip = f' title="Formula (year 2): {E(fm)}"' if fm else ""
        cells = "".join(f"<td>{dual(v[f'{c}{r}'].value, kind, scale, src_unit)}</td>" for c in cols)
        rows.append(f'<tr data-comment-target{" class=b" if bold else ""}><th scope="row"{tip}>{E(lab)}</th><td class="u">{E(unit)}</td>{cells}</tr>')
    th = "".join(f'<th scope="col">{h}</th>' for h in heads)
    return f'<div class="tw"><table class="an"><thead><tr><th class="c0" scope="col">Line</th><th class="u" scope="col">Unit</th>{th}</tr></thead><tbody>{"".join(rows)}</tbody></table></div>'


# ------------------------------------------------------------------ assumptions
def amt(x, kind, unit):
    """Input value with its other-currency equivalent: euros for FCFA and USD, FCFA for euro inputs."""
    u = unit or ""
    if u.startswith("€") and isinstance(x, (int, float)) and x:
        return f'{fmt_num(x, kind)}<span class="eur">{fcfa_fmt(x * FX)}</span>'
    return dual(x, kind, 1, u)


def assumptions_html():
    v, f = wv["Assumptions"], wf["Assumptions"]
    out, rows, mode = [], [], None

    def flush():
        if rows:
            if mode == "scalar":
                th = "<th scope='col'>Value used</th><th scope='col'>Prudent</th><th scope='col'>Base</th><th scope='col'>Ambitious</th>"
            else:
                th = "".join(f"<th scope='col'>Year {k}</th>" for k in range(1, 6))
            out.append(f'<div class="tw"><table class="as"><thead><tr><th class="c0" scope="col">Assumption</th><th class="u" scope="col">Unit</th>{th}<th class="note" scope="col">Source / note</th></tr></thead><tbody>{"".join(rows)}</tbody></table></div>')
            rows.clear()

    for r in range(6, v.max_row + 1):
        b, c3, d = v.cell(r, 2).value, v.cell(r, 3).value, v.cell(r, 4).value
        if b is None and d is None:
            continue
        if c3 == "Unit":
            flush()
            mode = "arr" if v.cell(r, 4).value == "Year 1" else "scalar"
            out.append(f'<h3 data-comment-target>{E(b)}</h3>')
            continue
        if d is None and v.cell(r, 9).value is None:
            rows.append(f'<tr class="sub"><th scope="rowgroup" colspan="8">{E(b)}</th></tr>')
            continue
        kind = kind_of(f.cell(r, 4).number_format)
        if kind == "date":
            dv = v.cell(r, 4).value
            val = dv.strftime("%d %b %Y") if hasattr(dv, "strftime") else E(dv)
        elif str(b).startswith("FCFA per"):
            val = fmt_num(d, "dec3" if d % 1 else "int")
        else:
            val = amt(d, kind, c3)
        note = E(v.cell(r, 9).value or "")
        if mode == "scalar":
            scen = [v.cell(r, k).value for k in (5, 6, 7)]
            has = any(s is not None for s in scen)
            sc = "".join(f"<td{' class=in' if has else ''}>{amt(s, kind, c3) if s is not None else ''}</td>" for s in scen)
            cells = f"<td class='{'used' if has else 'in'}'>{val}</td>{sc}"
        else:
            vals = [v.cell(r, k).value for k in range(4, 9)]
            cells = "".join(f"<td class='in'>{amt(x, kind, c3) if x is not None else ''}</td>" for x in vals)
            if all(x is None for x in vals[1:]):
                cells = f"<td class='in'>{val}</td>" + "<td></td>" * 4
        rows.append(f'<tr data-comment-target><th scope="row">{E(b)}</th><td class="u">{E(c3 or "")}</td>{cells}<td class="note">{note}</td></tr>')
    flush()
    return "".join(out)


def hiring_html():
    v = wv["Hiring plan"]
    rows = []
    for r in range(5, v.max_row + 1):
        role = v.cell(r, 3).value
        if not role:
            continue
        cells = [v.cell(r, k).value for k in range(1, 13)]
        n, dep, role, cty, cat, hc, st, en, sal, mgr, cost, note = cells
        if not hc:
            continue
        rows.append(f'<tr data-comment-target><td class="u">{n}</td><td>{E(dep)}</td><th scope="row">{E(role)}</th><td>{E(cty)}</td><td>{E(cat or "")}</td>'
                    f'<td>{fmt_num(hc, "int")}</td><td>{st}</td><td>{en or ""}</td><td>{E(mgr or "")}</td><td>{dual(sal, "int", 1, "FCFA")}</td><td>{dual(cost, "int", 1, "FCFA")}</td><td class="note">{E(note or "")}</td></tr>')
    th = "".join(f'<th scope="col">{h}</th>' for h in ["#", "Department", "Role", "Country", "Category", "People", "Start month", "End month", "Agency hire (€800)", "Gross salary (FCFA / month)", "Employer cost per person (FCFA / month)", "Note"])
    return f'<div class="tw"><table class="hp"><thead><tr>{th}</tr></thead><tbody>{"".join(rows)}</tbody></table></div>'


def blocks_html(sheet):
    """Render a sheet made of titled blocks (title row, optional italic subtitle, dark header row, data rows)."""
    v, f = wv[sheet], wf[sheet]
    out, rows, cols = [], [], []

    def flush():
        if rows:
            th = "".join(f'<th scope="col"{" class=c0" if k == 0 else ""}>{E(h)}</th>' for k, (c, h) in enumerate(cols))
            body = "".join(x if isinstance(x, str) else render(x) for x in rows)
            out.append(f'<div class="tw"><table class="an uc"><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table></div>')
            rows.clear()

    def render(x):
        tds = "".join((f'<th scope="row">{h}</th>' if k == 0 else f'<td{x["cls"][k]}>{h}</td>') for k, h in enumerate(x["cells"]))
        return f'<tr data-comment-target{" class=b" if x["bold"] else ""}>{tds}</tr>'

    def cell(r, c):
        x = v.cell(r, c).value
        fm = f.cell(r, c).number_format or ""
        if x is None or x == "":
            return ""
        if isinstance(x, str):
            return E(x)
        if "€" in fm:
            if abs(x) < 0.005:
                return '<span class="z">–</span>'
            t = f"{abs(x):,.2f}" if abs(x) < 1000 else f"{abs(x):,.0f}"
            t = "€" + t.replace(",", NNBSP)
            return f'<span class="neg">{MINUS}{t}</span>' if x < 0 else t
        return fmt_num(x, kind_of(fm))

    prev_title = False
    for r in range(4, v.max_row + 1):
        a = v.cell(r, 1)
        font, fill = a.font, a.fill
        if isinstance(a.value, str) and a.value.strip() == "in €":
            last = next((x for x in reversed(rows) if not isinstance(x, str)), None)
            if last is not None:
                for k, (c, h) in enumerate(cols):
                    e = v.cell(r, c).value
                    if isinstance(e, (int, float)) and abs(e) >= 0.005:
                        t = f"{abs(e):,.2f}" if abs(e) < 1000 else f"{abs(e):,.0f}"
                        t = f"{MINUS if e < 0 else ''}€" + t.replace(",", NNBSP)
                        last["cells"][k] += f'<span class="eur">{t}</span>'
            continue
        is_head = fill is not None and str(fill.fgColor.rgb).endswith("1F2A44")
        if is_head:
            cols = [(c, v.cell(r, c).value) for c in range(1, 16) if v.cell(r, c).value not in (None, "")]
            prev_title = False
            continue
        if a.value is None and all(v.cell(r, c).value is None for c in range(2, 16)):
            continue
        if font is not None and font.sz and float(font.sz) >= 12:
            flush()
            out.append(f'<h3 data-comment-target>{E(a.value)}</h3>')
            prev_title = True
            continue
        if font is not None and font.i and all(v.cell(r, c).value is None for c in range(2, 16)):
            if prev_title:
                out.append(f'<p class="lead">{E(a.value)}</p>')
            else:
                flush()
                out.append(f'<p class="lead small">{E(a.value)}</p>')
            prev_title = False
            continue
        prev_title = False
        if fill is not None and str(fill.fgColor.rgb).endswith("DCE6F2") and all(v.cell(r, c).value is None for c in range(2, 16)):
            rows.append(f'<tr class="sec"><th scope="rowgroup" colspan="{len(cols)}">{E(a.value)}</th></tr>')
            continue
        bold = bool(font and font.b)
        cls = [(" class=tx" if c == 2 and isinstance(v.cell(r, c).value, str) and len(v.cell(r, c).value) > 16 else (" class=u" if isinstance(v.cell(r, c).value, str) else ""))
               for c, h in cols]
        rows.append({"cells": [cell(r, c) for c, h in cols], "cls": cls, "bold": bold})
    flush()
    return "".join(out)


def regulatory_html():
    """Stage cards from the 'Regulatory path' sheet: stage rows have a name in A, requirement rows a topic in C."""
    v = wv["Regulatory path"]
    out, rows, intro = [], [], []

    def flush():
        if rows:
            out.append('<div class="rgw"><table class="rg"><colgroup><col class="k1"><col class="k2"><col class="k3"><col class="k4"><col class="k5"></colgroup>'
                       '<thead><tr><th scope="col">Topic</th><th scope="col">Requirement</th>'
                       '<th scope="col">Authority and legal basis</th><th scope="col">In the model</th><th scope="col">Status</th></tr></thead>'
                       f'<tbody>{"".join(rows)}</tbody></table></div>')
            rows.clear()

    for r in range(4, v.max_row + 1):
        a, b, c, d, e, f_, g = (v.cell(r, k).value for k in range(1, 8))
        if a == "Stage":
            continue
        if a and str(a).startswith("Note:"):
            intro.append(f'<p class="lead">{E(a[5:].strip())}</p>')
            continue
        if a:
            flush()
            out.append(f'<div class="stage" data-comment-target><h3>{E(a)}</h3><div class="when">{E(b or "")}</div><p>{E(d or "")}</p></div>')
            continue
        if c:
            st = str(g or "")
            cls = "ok" if st.lower().startswith("verified") else "tbc"
            rows.append(f'<tr data-comment-target><th scope="row">{E(c)}</th><td data-l="Requirement">{E(d or "")}</td><td class="src" data-l="Authority and legal basis">{E(e or "")}</td>'
                        f'<td data-l="In the model">{E(f_ or "")}</td><td data-l="Status"><span class="pill {cls}">{E(st)}</span></td></tr>')
    flush()
    return "".join(intro) + "".join(out)


# ------------------------------------------------------------------ headline + scenarios
def dash_values(path):
    w = load_workbook(path, data_only=True)
    D, Y = w["Dashboard"], w["P&L annual"]
    lab = {Y.cell(r, 1).value: r for r in range(1, Y.max_row + 1)}
    yv = lambda name, col: Y[f"{col}{lab[name]}"].value
    return {
        "scenario": D["B4"].value, "be": D["B7"].value, "be_date": D["B8"].value, "peak": D["B9"].value, "equity": D["B11"].value,
        "low": D["B13"].value, "low_m": D["B14"].value, "end_cash": D["B15"].value,
        "rev": [yv("Total revenue", c) for c in "DEFGH"], "ebitda": [yv("EBITDA", c) for c in "DEFGH"],
        "clients": [yv("Active clients", c) for c in "DEFGH"], "hc": [yv("Headcount", c) for c in "DEFGH"],
        "ni": [yv("Net income", c) for c in "DEFGH"], "cash": [yv("Closing cash (end of year)", c) for c in "DEFGH"],
    }


base = dash_values(XLSX)
s1, s3 = dash_values(SC1), dash_values(SC3)
fy0 = wv["P&L annual"]["D5"].value
M = lambda x: fmt_num(x / 1e6, "dec1") if isinstance(x, (int, float)) else E(x)
EURM = lambda x: fmt_num(x / FX / 1e6, "dec1") if isinstance(x, (int, float)) else E(x)


def be_text(d):
    if isinstance(d["be"], (int, float)):
        return f"Month {int(d['be'])} ({d['be_date'].strftime('%b %Y')})"
    return "Not within 5 years"


both = lambda x: f'{fcfa_fmt(x)}<span class="eur">{eur_fmt(x / FX)}</span>'
kpis = [
    ("EBITDA turns positive", be_text(base), "First month with positive EBITDA"),
    ("Peak funding need", eur_fmt(base['peak'] / FX), f"{fcfa_fmt(base['peak'])} of cumulative burn before break-even"),
    ("Equity in the plan", eur_fmt(base['equity'] / FX), f"{fcfa_fmt(base['equity'])}: seed, Series A and Series B"),
    ("Year 5 revenue", eur_fmt(base['rev'][4] / FX), f"{fcfa_fmt(base['rev'][4])} · {fmt_num(base['clients'][4], 'int')} active clients at the end of year 5"),
    ("Year 5 EBITDA", eur_fmt(base['ebitda'][4] / FX), f"{fcfa_fmt(base['ebitda'][4])} · {fmt_num(base['ebitda'][4] / base['rev'][4], 'pct')} of revenue"),
]
kpi_html = "".join(f'<div class="kpi" data-comment-target><div class="kl">{E(a)}</div><div class="kv">{b}</div><div class="ks">{c}</div></div>' for a, b, c in kpis)

scen_rows = [
    ("EBITDA turns positive", lambda d: be_text(d)),
    ("Peak funding need", lambda d: both(d["peak"])),
    ("Active clients, end of year 5", lambda d: fmt_num(d["clients"][4], "int")),
    ("Revenue year 3", lambda d: both(d["rev"][2])),
    ("Revenue year 5", lambda d: both(d["rev"][4])),
    ("EBITDA year 5", lambda d: both(d["ebitda"][4])),
    ("Lowest closing cash with the planned rounds", lambda d: both(d["low"])),
]
scen_html = '<div class="tw"><table class="an sc"><thead><tr><th class="c0" scope="col">Result</th><th scope="col">Prudent</th><th scope="col" class="hl">Base</th><th scope="col">Ambitious</th></tr></thead><tbody>' + "".join(
    f'<tr data-comment-target><th scope="row">{E(l)}</th><td>{fn(s1)}</td><td class="hl">{fn(base)}</td><td>{fn(s3)}</td></tr>' for l, fn in scen_rows) + "</tbody></table></div>"


# ------------------------------------------------------------------ charts (inline SVG, drawn to scale)
def bar_chart():
    rev = [x / FX / 1e6 for x in base["rev"]]
    eb = [x / FX / 1e6 for x in base["ebitda"]]
    hi = max(rev + [0]) * 1.1
    lo = min(eb + [0]) * 1.15
    W, H, L, R_, T, B = 640, 280, 48, 12, 16, 34
    ph = H - T - B
    y = lambda val: T + (hi - val) / (hi - lo) * ph
    step = 2 if hi - lo < 16 else 5
    ticks = [t for t in range(int(lo // step) * step, int(hi) + step, step) if lo <= t <= hi]
    g = []
    for t in ticks:
        g.append(f'<line x1="{L}" x2="{W - R_}" y1="{y(t):.1f}" y2="{y(t):.1f}" class="{"zero" if t == 0 else "grid"}"/>'
                 f'<text x="{L - 6}" y="{y(t) + 4:.1f}" class="ax" text-anchor="end">{t}</text>')
    gw = (W - L - R_) / 5
    bw = gw * 0.3
    for k in range(5):
        x0 = L + k * gw + gw * 0.18
        for j, (val, cls) in enumerate(((rev[k], "b1"), (eb[k], "b2"))):
            x = x0 + j * (bw + 4)
            top, bot = (y(val), y(0)) if val >= 0 else (y(0), y(val))
            g.append(f'<rect x="{x:.1f}" y="{top:.1f}" width="{bw:.1f}" height="{max(bot - top, 0.5):.1f}" class="{cls}" rx="2"><title>{"Revenue" if j == 0 else "EBITDA"} year {k + 1}: €{val:.2f}M ({val * FX / 1000:,.2f} bn FCFA)</title></rect>')
        g.append(f'<text x="{L + k * gw + gw / 2:.1f}" y="{H - 12}" class="ax" text-anchor="middle">Y{k + 1} · {fy0 + k}</text>')
    return (f'<figure class="chart" data-comment-target><figcaption>Revenue and EBITDA by year, € millions</figcaption>'
            f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="Revenue and EBITDA by year">{"".join(g)}</svg>'
            f'<div class="legend"><span><i class="b1"></i>Revenue</span><span><i class="b2"></i>EBITDA</span></div></figure>')


def cash_chart():
    Q = wv["Cash flow"]
    lab = {Q.cell(r, 1).value: r for r in range(1, Q.max_row + 1)}
    vals = [Q.cell(lab["Closing cash"], c).value / FX / 1e6 for c in MONTH_COLS]
    hi, lo = max(vals) * 1.1, min(0, min(vals)) * 1.1
    W, H, L, R_, T, B = 640, 240, 48, 12, 16, 30
    ph = H - T - B
    y = lambda v_: T + (hi - v_) / (hi - lo) * ph
    x = lambda i: L + i * (W - L - R_) / 59
    step = 1
    while (hi - lo) / step > 6:
        step *= 2
    ticks = [t for t in range(int(lo // step) * step, int(hi) + step, step) if lo <= t <= hi]
    g = [f'<line x1="{L}" x2="{W - R_}" y1="{y(t):.1f}" y2="{y(t):.1f}" class="{"zero" if t == 0 else "grid"}"/><text x="{L - 6}" y="{y(t) + 4:.1f}" class="ax" text-anchor="end">{t}</text>' for t in ticks]
    pts = " ".join(f"{x(i):.1f},{y(v_):.1f}" for i, v_ in enumerate(vals))
    g.append(f'<polygon points="{x(0):.1f},{y(0):.1f} {pts} {x(59):.1f},{y(0):.1f}" class="area"/>')
    g.append(f'<polyline points="{pts}" class="line"/>')
    for k in range(5):
        g.append(f'<text x="{x(k * 12 + 6):.1f}" y="{H - 10}" class="ax" text-anchor="middle">Y{k + 1}</text>')
    lowi = vals.index(min(vals))
    g.append(f'<circle cx="{x(lowi):.1f}" cy="{y(vals[lowi]):.1f}" r="4" class="dot"><title>Lowest: €{vals[lowi]:.2f}M ({vals[lowi] * FX:,.0f}M FCFA) in month {lowi + 1}</title></circle>')
    return (f'<figure class="chart" data-comment-target><figcaption>Closing cash by month, € millions (after the funding rounds)</figcaption>'
            f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="Closing cash by month">{"".join(g)}</svg></figure>')


# ------------------------------------------------------------------ page
tabs = [
    ("overview", "Overview"), ("assumptions", "Assumptions"), ("unitcosts", "Unit costs"), ("hiring", "Hiring plan"), ("pnl", "P&L by year"), ("pnlm", "P&L by month"),
    ("cash", "Cash flow"), ("costs", "Costs"), ("sales", "Sales & marketing"), ("revenue", "Clients & revenue"), ("credit", "Credit"),
    ("staff", "Staff"), ("unit", "Unit economics"), ("market", "Market"), ("regulatory", "Regulatory path"),
]
sections = {
    "overview": ("Overview", f"Base scenario. Five years from {nv('Start_Date'):%B %Y}, Senegal first, Côte d'Ivoire from month {nv('Launch_CI')}. Partners hold the licences until Tabax buys a microfinance company (SFD) in month {nv('SFD_Acq_Month')}.",
                 f'<div class="kpis">{kpi_html}</div><div class="charts">{bar_chart()}{cash_chart()}</div>'
                 f'<h3 data-comment-target>Three scenarios</h3><p class="lead">The scenario switch changes sales productivity, partner and digital acquisition, churn, interchange, the bank income share on balances, credit volumes and credit losses.</p>{scen_html}'
                 f'<h3 data-comment-target>Year by year (FCFA millions, euros underneath)</h3>' + annual_table("Dashboard", cols="DEFGH", head_row=19, first=20, stop_label="Checks", scale_fcfa={"div": 1e6, "unit": "M FCFA"})),
    "assumptions": ("Assumptions", "Every input of the model with its source. Values in the 'Value used' column are those of the base scenario.", assumptions_html()),
    "unitcosts": ("Unit costs", "What one of each thing costs or earns: one employee by role, one business client by plan, one loan, and the price of every fixed item. Every figure comes from the inputs on Assumptions and Hiring plan.", blocks_html("Unit costs")),
    "hiring": ("Hiring plan", f"The fixed roles. Employer cost = gross salary + {nv('Employer_Charges'):.0%} charges. Customer support, developers, field sales reps and credit analysts grow with clients and the loan book: see the Staff tab.", hiring_html()),
    "pnl": ("P&L by year", "FCFA millions, with the euro equivalent under each amount.", annual_table("P&L annual", scale_fcfa={"div": 1e6, "unit": "M FCFA"})),
    "pnlm": ("P&L by month", "Thousands of FCFA (k FCFA), with the euro equivalent under each amount. Hover a line name to see its formula.", ts_table("P&L monthly")),
    "cash": ("Cash flow", "Thousands of FCFA, euros underneath. Cash set aside (card collateral, credit reserves, equity in the own SFD's loans) is excluded from closing cash.", ts_table("Cash flow")),
    "costs": ("Costs", "Every cost line in thousands of FCFA, euros underneath: variable costs, then fixed costs by category and the 5% contingency.", ts_table("Costs")),
    "sales": ("Sales & marketing", "New clients by channel and country, and the non-staff marketing budget (k FCFA).", ts_table("Sales & marketing")),
    "revenue": ("Clients & revenue", "Active clients by plan, usage and revenue lines except credit (k FCFA).", ts_table("Clients & revenue")),
    "credit": ("Credit", f"Loan book by product and credit income (k FCFA). No loans before month {min(nv('Start_Factoring'), nv('Start_Bridge'))}. A partner SFD lends until month {nv('Own_Funding_Month') - 1}, then Tabax's own SFD in Senegal; Côte d'Ivoire loans stay with a partner.", ts_table("Credit")),
    "staff": ("Staff", "The team that grows with volume (support per client, developers per client, field reps per client in each country, credit analysts per loan book), then headcount and employer cost by department (k FCFA).", ts_table("Staff", from_label="Team that grows with volume")),
    "unit": ("Unit economics", "Per-client economics and efficiency ratios by year.", annual_table("Unit economics", cols="DEFGH", first=6)),
    "regulatory": ("Regulatory path", "What each stage of the plan needs from regulators in the WAEMU (BCEAO zone), Senegal and Côte d'Ivoire, and where its cost sits in the model.", regulatory_html()),
    "market": ("Market", "Formal employers in Senegal and Côte d'Ivoire, the target market and the share the plan reaches.", annual_table("Market", cols="DEFGH", head_row=4, first=5)),
}
nav = "".join(f'<button type="button" role="tab" id="t-{k}" aria-controls="p-{k}" aria-selected="{"true" if i == 0 else "false"}" data-tab="{k}">{E(lab)}</button>' for i, (k, lab) in enumerate(tabs))
panels = "".join(
    f'<section class="panel" id="p-{k}" role="tabpanel" aria-labelledby="t-{k}"{"" if i == 0 else " hidden"}>'
    f'<div class="ph"><div><h2>{E(sections[k][0])}</h2><p class="lead">{sections[k][1]}</p></div>'
    f'<button type="button" class="cbtn" data-comment-for="p-{k}" hidden>Comment on this tab</button></div>{sections[k][2]}</section>'
    for i, (k, lab) in enumerate(tabs))
xlsx_b64 = base64.b64encode(open(XLSX, "rb").read()).decode()
today = datetime.date.today().strftime("%d %B %Y").lstrip("0")

page = f"""<title>Tabax Business Plan</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Epilogue:wght@600;800&family=IBM+Plex+Sans:wght@400;500;600&display=swap">
<style>
/* Layout: a working financial file. Header with the five numbers that matter, tabs mirroring the Excel sheets, wide tables that scroll inside their own frame with a sticky first column. */
:root {{
  --bg: #F6F7F9; --surface: #FFFFFF; --ink: #161616; --muted: #5A6170; --line: #E3E6EB; --soft: #EEF1F5;
  --accent: #6FB6EE; --accent-ink: #1C5F96; --neg: #B42318; --hl: #E8F3FC;
  --ok: #1E7F3C; --okbg: #DDF0E1; --tbc: #8A5A00; --tbcbg: #FFF1D6;
  --display: 'Epilogue', 'Helvetica Neue', Arial, sans-serif; --body: 'IBM Plex Sans', 'Helvetica Neue', Arial, sans-serif;
}}
@media (prefers-color-scheme: dark) {{ :root:not([data-theme="light"]) {{
  --bg: #0F1216; --surface: #161A20; --ink: #ECEFF3; --muted: #9AA3AF; --line: #2A3038; --soft: #1C2128;
  --accent: #6FB6EE; --accent-ink: #8FC9F5; --neg: #FF8A80; --hl: #17283A;
  --ok: #8FD6A3; --okbg: #17301F; --tbc: #F5C76B; --tbcbg: #3A2D12; color-scheme: dark; }} }}
:root[data-theme="dark"] {{
  --bg: #0F1216; --surface: #161A20; --ink: #ECEFF3; --muted: #9AA3AF; --line: #2A3038; --soft: #1C2128;
  --accent: #6FB6EE; --accent-ink: #8FC9F5; --neg: #FF8A80; --hl: #17283A;
  --ok: #8FD6A3; --okbg: #17301F; --tbc: #F5C76B; --tbcbg: #3A2D12; color-scheme: dark; }}
body {{ background: var(--bg); color: var(--ink); font-family: var(--body); font-size: 14px; line-height: 1.45; }}
.wrap {{ max-width: 1240px; margin: 0 auto; padding-inline: 16px; padding-block: 24px 64px; }}
header.top {{ display: flex; flex-wrap: wrap; gap: 16px 24px; align-items: flex-end; justify-content: space-between; }}
.eyebrow {{ font-size: 12px; font-weight: 600; letter-spacing: .12em; text-transform: uppercase; color: var(--muted); display: flex; align-items: center; gap: 8px; }}
.eyebrow i {{ width: 10px; height: 10px; border-radius: 3px; background: var(--accent); box-shadow: inset 0 0 0 2px var(--ink); }}
h1 {{ font-family: var(--display); font-weight: 800; font-size: clamp(26px, 4vw, 38px); letter-spacing: -.02em; margin: 6px 0 4px; text-wrap: balance; }}
.sub {{ color: var(--muted); margin: 0; max-width: 70ch; }}
.actions {{ display: flex; gap: 8px; flex-wrap: wrap; }}
.btn {{ font: inherit; font-weight: 600; border: 1.5px solid var(--ink); background: var(--accent); color: #161616; border-radius: 20px; height: 38px; padding: 0 16px; cursor: pointer; }}
.btn.ghost {{ background: var(--surface); color: var(--ink); border-color: var(--line); }}
.btn:focus-visible, nav button:focus-visible, .cbtn:focus-visible {{ outline: 3px solid var(--accent); outline-offset: 2px; }}
.status {{ font-size: 13px; color: var(--muted); min-height: 1.2em; }}
nav.tabs {{ position: sticky; top: env(safe-area-inset-top, 0px); z-index: 5; background: var(--bg); margin-top: 20px; padding-block: 8px; border-bottom: 1px solid var(--line); display: flex; flex-wrap: wrap; gap: 4px; }}
@media (max-width: 760px) {{ nav.tabs {{ flex-wrap: nowrap; overflow-x: auto; scrollbar-width: thin; }} }}
nav.tabs button {{ font: inherit; font-weight: 600; font-size: 13px; color: var(--muted); background: none; border: 0; border-radius: 16px; padding: 7px 12px; white-space: nowrap; cursor: pointer; }}
nav.tabs button[aria-selected="true"] {{ background: var(--ink); color: var(--bg); }}
.panel {{ margin-top: 20px; }}
.ph {{ display: flex; flex-wrap: wrap; gap: 12px; justify-content: space-between; align-items: flex-start; }}
.ph > div {{ min-width: 0; }}
h2 {{ font-family: var(--display); font-weight: 800; font-size: 24px; margin: 0; letter-spacing: -.01em; }}
h3 {{ font-family: var(--display); font-weight: 600; font-size: 17px; margin: 28px 0 8px; }}
.lead {{ color: var(--muted); margin: 4px 0 12px; max-width: 80ch; }}
.cbtn {{ font: inherit; font-size: 13px; font-weight: 600; color: var(--accent-ink); background: var(--surface); border: 1px solid var(--line); border-radius: 16px; padding: 6px 12px; cursor: pointer; }}
.kpis {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 10px; margin-block: 8px 16px; }}
.kpi {{ background: var(--surface); border: 1px solid var(--line); border-radius: 14px; padding: 12px 14px; min-width: 0; }}
.kl {{ font-size: 12px; font-weight: 600; letter-spacing: .06em; text-transform: uppercase; color: var(--muted); }}
.kv {{ font-family: var(--display); font-weight: 800; font-size: 24px; margin-top: 4px; font-variant-numeric: tabular-nums; }}
.ks {{ font-size: 12px; color: var(--muted); }}
.charts {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(300px, 1fr)); gap: 12px; }}
.chart {{ margin: 0; background: var(--surface); border: 1px solid var(--line); border-radius: 14px; padding: 12px 14px; min-width: 0; }}
.chart figcaption {{ font-weight: 600; font-size: 13px; margin-bottom: 6px; }}
.chart svg {{ width: 100%; height: auto; display: block; }}
.chart .grid {{ stroke: var(--line); stroke-width: 1; }}
.chart .zero {{ stroke: var(--muted); stroke-width: 1.2; }}
.chart .ax {{ fill: var(--muted); font-size: 11px; font-family: var(--body); }}
.chart .b1 {{ fill: var(--accent); }}
.chart .b2 {{ fill: var(--ink); }}
.chart .line {{ fill: none; stroke: var(--accent-ink); stroke-width: 2; }}
.chart .area {{ fill: var(--hl); }}
.chart .dot {{ fill: var(--neg); }}
.legend {{ display: flex; gap: 14px; font-size: 12px; color: var(--muted); margin-top: 4px; }}
.legend i {{ display: inline-block; width: 10px; height: 10px; border-radius: 2px; margin-right: 6px; vertical-align: -1px; }}
.legend i.b1 {{ background: var(--accent); }} .legend i.b2 {{ background: var(--ink); }}
.tw {{ overflow-x: auto; border: 1px solid var(--line); border-radius: 12px; background: var(--surface); max-width: 100%; }}
.tw + .tw {{ margin-top: 12px; }}
table {{ border-collapse: separate; border-spacing: 0; font-variant-numeric: tabular-nums; width: max-content; min-width: 100%; }}
th, td {{ padding: 6px 10px; border-bottom: 1px solid var(--line); text-align: right; white-space: nowrap; font-weight: 400; }}
thead th {{ position: sticky; top: 0; background: var(--soft); font-size: 12px; font-weight: 600; color: var(--muted); z-index: 1; }}
th[scope="row"], th.c0 {{ text-align: left; position: sticky; left: 0; background: var(--surface); z-index: 2; min-width: 240px; max-width: 360px; white-space: normal; }}
thead th.c0 {{ background: var(--soft); z-index: 3; }}
th[scope="row"][title] {{ cursor: help; text-decoration: underline dotted var(--line); text-underline-offset: 3px; }}
td.u, th.u {{ text-align: left; color: var(--muted); font-size: 12px; }}
td.t, th.t {{ font-weight: 600; background: var(--soft); }}
tr.b th, tr.b td {{ font-weight: 600; }}
tr.b th[scope="row"] {{ font-weight: 700; }}
tr.sec th {{ text-align: left; background: var(--soft); font-weight: 700; color: var(--ink); font-size: 12px; letter-spacing: .04em; text-transform: uppercase; position: static; }}
tr.sec th {{ position: sticky; left: 0; }}
tr.sub th {{ text-align: left; font-weight: 700; padding-top: 12px; }}
th.yr {{ color: var(--ink); }}
.fy {{ font-weight: 400; color: var(--muted); }}
.neg {{ color: var(--neg); }}
.eur {{ display: block; font-size: 11px; color: var(--muted); font-weight: 400; line-height: 1.25; }}
.kv .eur {{ display: none; }}
.kpi .ks {{ margin-top: 2px; }}
.z {{ color: var(--muted); }}
td.in {{ color: var(--accent-ink); }}
td.used {{ font-weight: 600; }}
td.note, th.note {{ text-align: left; white-space: normal; min-width: 280px; max-width: 520px; color: var(--muted); font-size: 12px; }}
table.as th[scope="row"] {{ min-width: 260px; }}
table.hp td {{ text-align: left; }}
table.hp th[scope="row"] {{ position: static; }}
.sc .hl {{ background: var(--hl); font-weight: 600; }}
tbody tr:hover td, tbody tr:hover th[scope="row"] {{ background: var(--hl); }}
.stage {{ margin-top: 28px; }}
table.uc td.tx {{ text-align: left; white-space: normal; min-width: 220px; max-width: 320px; color: var(--muted); font-size: 12px; }}
.lead.small {{ font-size: 12px; margin-top: 6px; }}
.stage h3 {{ margin: 0; }}
.stage .when {{ font-size: 12px; font-weight: 600; letter-spacing: .06em; text-transform: uppercase; color: var(--accent-ink); margin-top: 2px; }}
.stage p {{ margin: 6px 0 10px; max-width: 85ch; }}
.rgw {{ border: 1px solid var(--line); border-radius: 12px; background: var(--surface); overflow: hidden; }}
table.rg {{ width: 100%; min-width: 0; table-layout: fixed; }}
table.rg col.k1 {{ width: 15%; }} table.rg col.k2 {{ width: 40%; }} table.rg col.k3 {{ width: 20%; }} table.rg col.k4 {{ width: 15%; }} table.rg col.k5 {{ width: 10%; }}
table.rg th, table.rg td {{ text-align: left; white-space: normal; vertical-align: top; overflow-wrap: anywhere; }}
table.rg thead th {{ position: static; }}
table.rg th[scope="row"] {{ position: static; min-width: 0; max-width: none; font-weight: 600; }}
table.rg td.src {{ color: var(--muted); font-size: 12px; }}
table.rg .pill {{ white-space: normal; }}
@media (max-width: 760px) {{
  table.rg, table.rg tbody, table.rg tr, table.rg th, table.rg td {{ display: block; width: auto; }}
  table.rg thead, table.rg colgroup {{ display: none; }}
  table.rg tr {{ padding: 10px 12px; border-bottom: 1px solid var(--line); }}
  table.rg th, table.rg td {{ border: 0; padding: 3px 0; }}
  table.rg td[data-l]::before {{ content: attr(data-l); display: block; font-size: 11px; font-weight: 600; letter-spacing: .04em; text-transform: uppercase; color: var(--muted); margin-top: 4px; }}
}}
.pill {{ display: inline-block; font-size: 11px; font-weight: 600; border-radius: 10px; padding: 2px 8px; white-space: nowrap; }}
.pill.ok {{ background: var(--okbg); color: var(--ok); }}
.pill.tbc {{ background: var(--tbcbg); color: var(--tbc); }}
footer {{ margin-top: 40px; color: var(--muted); font-size: 12px; max-width: 90ch; }}
@media (max-width: 600px) {{ th[scope="row"], th.c0 {{ min-width: 170px; max-width: 200px; }} .kv {{ font-size: 21px; }} }}
</style>
<div class="wrap">
<header class="top">
  <div>
    <div class="eyebrow"><i></i>Tabax (working name) · business plan · base scenario</div>
    <h1>Tabax business plan</h1>
    <p class="sub">B2B neobank for SMEs in Senegal, then Côte d'Ivoire. Five years, month by month, from the Excel model of {today}. Comment on any line: Claude updates the Excel formulas and republishes this page.</p>
  </div>
  <div>
    <div class="actions"><button type="button" class="btn" id="dl" hidden>Download the Excel model</button></div>
    <div class="status" id="dl-status" role="status"></div>
  </div>
</header>
<nav class="tabs" role="tablist" aria-label="Sheets">{nav}</nav>
{panels}
<footer>Source: the Excel model in this project (every number is a formula of the Assumptions and Hiring plan sheets). Amounts exclude VAT and the 17% tax on financial activities. Prices and costs marked as estimates on the Assumptions tab must be replaced with partner quotes.</footer>
</div>
<script type="application/octet-stream" id="xlsx">{xlsx_b64}</script>
<script>
(function () {{
  var tabs = Array.prototype.slice.call(document.querySelectorAll('nav.tabs button'));
  function show(key, focus) {{
    var found = false;
    tabs.forEach(function (b) {{
      var on = b.getAttribute('data-tab') === key;
      if (on) found = true;
      b.setAttribute('aria-selected', on ? 'true' : 'false');
      document.getElementById('p-' + b.getAttribute('data-tab')).hidden = !on;
      if (on) {{ var n = b.parentNode; if (n.scrollWidth > n.clientWidth) n.scrollLeft = Math.max(0, b.offsetLeft - n.offsetLeft - 16); }}
    }});
    if (!found) show('overview');
  }}
  tabs.forEach(function (b) {{
    b.addEventListener('click', function () {{ show(b.getAttribute('data-tab'), true); try {{ history.replaceState(null, '', '#' + b.getAttribute('data-tab')); }} catch (e) {{}} }});
  }});
  var h = (location.hash || '').slice(1);
  if (h) show(h);

  function start() {{
    if (!window.claude || !window.claude.use) return;
    window.claude.use('downloads').then(function (dl) {{
      if (!dl) return;
      var btn = document.getElementById('dl'), st = document.getElementById('dl-status');
      btn.hidden = false;
      btn.addEventListener('click', function () {{
        var raw = atob(document.getElementById('xlsx').textContent.trim());
        var bytes = new Uint8Array(raw.length);
        for (var i = 0; i < raw.length; i++) bytes[i] = raw.charCodeAt(i);
        st.textContent = '';
        dl.save({{filename: 'Tabax business plan.xlsx', data: new Blob([bytes])}}).then(function () {{
          st.textContent = 'Excel model saved.';
        }}, function (err) {{
          var code = err && err.code;
          if (code === 'declined') st.textContent = 'Download cancelled.';
          else if (code === 'rate_limited') st.textContent = 'A download prompt is already open.';
          else {{ st.textContent = 'Download is not available here. Ask Claude for the file.'; btn.hidden = true; }}
        }});
      }});
    }});
    window.claude.use('comments').then(function (cm) {{
      if (!cm) return;
      Array.prototype.forEach.call(document.querySelectorAll('.cbtn'), function (b) {{
        b.hidden = false;
        b.addEventListener('click', function () {{
          var el = document.getElementById(b.getAttribute('data-comment-for'));
          cm.openComposer({{element: el.querySelector('h2') || el}}).catch(function (err) {{
            if (err && (err.code === 'unavailable' || err.code === 'not_granted')) {{
              Array.prototype.forEach.call(document.querySelectorAll('.cbtn'), function (x) {{ x.hidden = true; }});
            }}
          }});
        }});
      }});
    }});
  }}
  start();
}})();
</script>
"""
open(OUT, "w").write(page)
print(round(len(page) / 1024), "KB")
